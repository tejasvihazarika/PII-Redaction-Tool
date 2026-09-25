import os
import uuid
import shutil
import logging
from pathlib import Path
from io import BytesIO
from typing import List, Dict, Any, Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
import docx

from redact_pii import detect_pii, apply_redaction, PIIRedactor, nlp

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("pii_server")

app = FastAPI(
    title="Redactly PII API",
    description="Refined SpaCy ML & Regex PII Detection & Redaction Service",
    version="1.0.0"
)

# Enable CORS for frontend integration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

OUTPUT_DIR = Path("./output")
OUTPUT_DIR.mkdir(exist_ok=True)

class ScanInput(BaseModel):
    text: str

class RedactInput(BaseModel):
    text: str
    findings: Optional[List[Dict[str, Any]]] = None
    mode: str = "fake"

class ExportTextInput(BaseModel):
    text: str
    findings: Optional[List[Dict[str, Any]]] = None
    mode: str = "fake"
    filename: str = "redacted-text.txt"

@app.get("/api/health")
async def health_check():
    return {
        "status": "online",
        "spacy_loaded": nlp is not None,
        "model": "en_core_web_sm" if nlp else "regex_fallback"
    }

@app.post("/api/scan")
async def scan_text(input_data: ScanInput):
    try:
        findings = detect_pii(input_data.text)
        return {"findings": findings, "count": len(findings)}
    except Exception as e:
        logger.error(f"Error scanning text for PII: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/redact")
async def redact_text(input_data: RedactInput):
    try:
        findings = input_data.findings if input_data.findings is not None else detect_pii(input_data.text)
        redacted = apply_redaction(input_data.text, findings, input_data.mode)
        return {"redacted_text": redacted, "findings": findings}
    except Exception as e:
        logger.error(f"Error redacting text: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/export-text")
async def export_text(input_data: ExportTextInput):
    try:
        findings = input_data.findings if input_data.findings is not None else detect_pii(input_data.text)
        redacted_text = apply_redaction(input_data.text, findings, input_data.mode)
        safe_name = Path(input_data.filename).name or "redacted-text.txt"
        if not safe_name.lower().endswith(".txt"):
            safe_name += ".txt"
        
        out_filename = f"{uuid.uuid4()[:8]}_{safe_name}"
        out_path = OUTPUT_DIR / out_filename
        out_path.write_text(redacted_text, encoding="utf-8")
        
        return {
            "download_url": f"/api/download-text/{out_filename}",
            "filename": safe_name,
            "redacted_text": redacted_text
        }
    except Exception as e:
        logger.error(f"Error exporting text: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/upload-docx")
async def upload_docx(file: UploadFile = File(...), mode: str = Form("fake")):
    try:
        file_bytes = await file.read()
        doc = docx.Document(BytesIO(file_bytes))
        
        all_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                all_text.append(p.text)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        all_text.append(cell.text)
        
        full_text = "\n".join(all_text)
        findings = detect_pii(full_text)
        
        # Apply redaction paragraph by paragraph & table cell by table cell
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                p_findings = detect_pii(paragraph.text)
                paragraph.text = apply_redaction(paragraph.text, p_findings, mode=mode)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        c_findings = detect_pii(cell.text)
                        cell.text = apply_redaction(cell.text, c_findings, mode=mode)

        original_filename = Path(file.filename or "document.docx").stem
        out_name = f"{original_filename}_redacted_{uuid.uuid4()[:6]}.docx"
        out_path = OUTPUT_DIR / out_name
        doc.save(str(out_path))

        return {
            "download_url": f"/api/download/{out_name}",
            "filename": out_name,
            "original_text": full_text,
            "findings": findings
        }
    except Exception as e:
        logger.error(f"Error processing docx upload: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/download/{filename}")
async def download_file(filename: str):
    file_path = OUTPUT_DIR / Path(filename).name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(
        path=file_path,
        filename=Path(filename).name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

@app.get("/api/download-text/{filename}")
async def download_text(filename: str):
    file_path = OUTPUT_DIR / Path(filename).name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(
        path=file_path,
        filename=Path(filename).name,
        media_type="text/plain; charset=utf-8"
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
